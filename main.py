import os
import io
import streamlit as st
from openai import OpenAI
from pypdf import PdfReader
from docx import Document

# Configurazione pagina
st.set_page_config(page_title="App di Numerologia Evolutiva", layout="centered")

# Impedisce traduzioni automatiche alterate dal browser
st.markdown("""
    <html lang="it" class="notranslate">
    <head><meta name="google" content="notranslate" /></head>
    <style>
        .element-container, .stMarkdown, p, h1, h2, h3 {
            font-family: 'Georgia', serif;
        }
    </style>
""", unsafe_allow_html=True)

def check_password():
    if "password_correct" not in st.session_state:
        st.session_state["password_correct"] = False
    if st.session_state["password_correct"]:
        return True

    st.title("🔒 Accesso Riservato")
    password_input = st.text_input("Inserisci la password per accedere all'app:", type="password")
    
    if st.button("Accedi"):
        if password_input == "numerologia2026":
            st.session_state["password_correct"] = True
            st.rerun()
        else:
            st.error("❌ Password errata!")
    return False

# --- TABELLA PITAGORICA ESATTA ---
TABELLA_PITAGORICA = {
    'a': 1, 'j': 1, 's': 1,
    'b': 2, 'k': 2, 't': 2,
    'c': 3, 'l': 3, 'u': 3,
    'd': 4, 'm': 4, 'v': 4,
    'e': 5, 'n': 5, 'w': 5,
    'f': 6, 'o': 6, 'x': 6,
    'g': 7, 'p': 7, 'y': 7,
    'h': 8, 'q': 8, 'z': 8,
    'i': 9, 'r': 9
}

VOCALI_BASE = set(['a', 'e', 'i', 'o', 'u'])

def riduci(n):
    while n > 9 and n not in [11, 22, 33, 44]:
        n = sum(int(c) for c in str(n))
    return n

def riduci_singola(n):
    while n > 9:
        n = sum(int(c) for c in str(n))
    return n

def calcola_segno_zodiacale(g, m):
    g, m = int(g), int(m)
    if (m == 3 and g >= 21) or (m == 4 and g <= 20): return "Ariete (Fuoco)"
    if (m == 4 and g >= 21) or (m == 5 and g <= 20): return "Toro (Terra)"
    if (m == 5 and g >= 21) or (m == 6 and g <= 21): return "Gemelli (Aria)"
    if (m == 6 and g >= 22) or (m == 7 and g <= 22): return "Cancro (Acqua)"
    if (m == 7 and g >= 23) or (m == 8 and g <= 23): return "Leone (Fuoco)"
    if (m == 8 and g >= 24) or (m == 9 and g <= 22): return "Vergine (Terra)"
    if (m == 9 and g >= 23) or (m == 10 and g <= 22): return "Bilancia (Aria)"
    if (m == 10 and g >= 23) or (m == 11 and g <= 21): return "Scorpione (Acqua)"
    if (m == 11 and g >= 22) or (m == 12 and g <= 21): return "Sagittario (Fuoco)"
    if (m == 12 and g >= 22) or (m == 1 and g <= 20): return "Capricorno (Terra)"
    if (m == 1 and g >= 21) or (m == 2 and g <= 19): return "Acquario (Aria)"
    if (m == 2 and g >= 20) or (m == 3 and g <= 20): return "Pesci (Acqua)"
    return "Non definito"

def calcola_nome_cognome(testo):
    testo_clean = testo.lower()
    vocali_trovate = []
    consonanti_trovate = []
    
    # Conteggio Piani dell'Essere (F.E.M.S.) basato sui numeri pitagorici
    fems_count = {'Fisico': 0, 'Emotivo': 0, 'Mentale': 0, 'Spirituale': 0}
    frequenza_numeri = {i: 0 for i in range(1, 10)}
    
    for char in testo_clean:
        if char in TABELLA_PITAGORICA:
            val = TABELLA_PITAGORICA[char]
            frequenza_numeri[val] += 1
            
            # Assegnazione corretta ai piani in base al valore numerico pitagorico
            if val in [1, 5]:
                fems_count['Mentale'] += 1
            elif val in [2, 3, 6, 9]:
                fems_count['Emotivo'] += 1
            elif val in [4, 8]:
                fems_count['Fisico'] += 1
            elif val == 7:
                fems_count['Spirituale'] += 1
            
            if char in VOCALI_BASE or char == 'y':
                vocali_trovate.append((char.upper(), val))
            else:
                consonanti_trovate.append((char.upper(), val))
                
    tot_v = sum(v[1] for v in vocali_trovate)
    tot_c = sum(c[1] for c in consonanti_trovate)
    tot_io = tot_v + tot_c
    
    return {
        "anima_tot": tot_v, "anima_rid": riduci(tot_v),
        "persona_tot": tot_c, "persona_rid": riduci(tot_c),
        "io_tot": tot_io, "io_rid": riduci(tot_io),
        "fems_count": fems_count,
        "frequenza_numeri": frequenza_numeri
    }

def calcola_data_nascita(g, m, a):
    g, m, a = int(g), int(m), int(a)
    dest_tot = g + m + a
    dest_rid = riduci(dest_tot)
    sole_tot = g + m
    sole_rid = riduci(sole_tot)
    eta_innesto = g + m
    
    origine_tot = sum(int(c) for c in str(a))
    origine_rid = riduci(origine_tot)
    dono_tot = int(str(a)[-2:])
    dono_rid = riduci(dono_tot)
    
    m_s = riduci_singola(m)
    g_s = riduci_singola(g)
    a_s = riduci_singola(sum(int(c) for c in str(a)))
    
    sfida_infanzia = abs(g_s - a_s)
    sfida_maturita = abs(m_s - g_s)
    sfida_principale = abs(sfida_maturita - sfida_infanzia)
    
    cifre_data = [int(c) for c in f"{g}{m}{a}" if c.isdigit() and c != '0']
    
    return {
        "dest_tot": dest_tot, "dest_rid": dest_rid,
        "approdo": g,
        "sole_tot": sole_tot, "sole_rid": sole_rid,
        "eta_innesto": eta_innesto,
        "origine_rid": origine_rid,
        "dono_rid": dono_rid,
        "sfida_infanzia": sfida_infanzia,
        "sfida_maturita": sfida_maturita,
        "sfida_principale": sfida_principale,
        "cifre_data": cifre_data
    }

def calcola_lezioni_e_dominanti(freq_nome, cifre_data, numeri_sintesi):
    freq_totale = freq_nome.copy()
    for c in cifre_data:
        if c in freq_totale:
            freq_totale[c] += 1
            
    numeri_sintesi_set = set(numeri_sintesi)
    
    # Lezioni Karmiche: numeri a frequenza 0 che NON compaiono nei numeri di sintesi (Croce Cardinale)
    lezioni_karmiche = [
        num for num, count in freq_totale.items() 
        if count == 0 and num not in numeri_sintesi_set
    ]
    
    archetipi_dominanti = [num for num, count in freq_totale.items() if count >= 3]
    
    return freq_totale, lezioni_karmiche, archetipi_dominanti

def crea_documento_word(nome_soggetto, dati_sintesi, testo_analisi):
    doc = Document()
    doc.add_heading("Mappa Numerologica Evolutiva", level=1)
    doc.add_heading(f"Profilo di: {nome_soggetto}", level=2)
    
    doc.add_heading("Quadro Numerico", level=3)
    for riga in dati_sintesi.split("\n"):
        if riga.strip():
            doc.add_paragraph(riga)
            
    doc.add_heading("Analisi Dettagliata", level=3)
    doc.add_paragraph(testo_analisi)
    
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream

if check_password():
    st.title("✨ Mappa Numerologica Evolutiva")
    st.write("Analisi narrativa con Piani F.E.M.S., Segno Zodiacale e Lezioni Karmiche corrette.")

    api_key = st.secrets.get("OPENAI_API_KEY") or os.environ.get("OPENAI_API_KEY")

    if not api_key:
        st.warning("⚠️ Inserisci la chiave API di OpenAI nei segreti di Streamlit.")
    else:
        client = OpenAI(api_key=api_key)

    uploaded_pdf = st.file_uploader("Carica la dispensa unica in PDF", type=["pdf"])

    pdf_text = ""
    if uploaded_pdf is not None:
        try:
            reader = PdfReader(uploaded_pdf)
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    pdf_text += text + "\n"
            st.success("PDF caricato ed estratto con successo!")
        except Exception as e:
            st.error(f"Errore nella lettura del PDF: {e}")

    col1, col2 = st.columns(2)
    with col1:
        nome_cognome = st.text_input("Nome e Cognome:", value="Giovanni Pittia")
    with col2:
        data_nascita = st.text_input("Data di Nascita (GG/MM/AAAA):", value="22/03/2009")

    if st.button("Genera Analisi Approfondita") and api_key:
        if nome_cognome and data_nascita:
            with st.spinner("Elaborazione matematica e stesura del profilo in corso..."):
                try:
                    res_nc = calcola_nome_cognome(nome_cognome)
                    p = data_nascita.split("/")
                    res_d = calcola_data_nascita(p[0], p[1], p[2])
                    
                    q_tot = res_nc["io_tot"] + res_d["dest_tot"]
                    q_rid = riduci(q_tot)
                    
                    segno_zodiacale = calcola_segno_zodiacale(p[0], p[1])
                    
                    numeri_cardinali = [
                        res_nc['anima_rid'], 
                        res_nc['persona_rid'], 
                        res_nc['io_rid'], 
                        res_d['dest_rid'], 
                        q_rid
                    ]
                    
                    freq_totale, lezioni_karmiche, archetipi_dominanti = calcola_lezioni_e_dominanti(
                        res_nc["frequenza_numeri"], 
                        res_d["cifre_data"],
                        numeri_cardinali
                    )
                    
                    # 1. QUADRO SINTETICO
                    st.markdown("---")
                    st.subheader("📊 Mappa dei Numeri Fondamentali")
                    c1, c2, c3, c4, c5 = st.columns(5)
                    c1.metric("Anima (A)", f"{res_nc['anima_tot']}/{res_nc['anima_rid']}")
                    c2.metric("Persona (P)", f"{res_nc['persona_tot']}/{res_nc['persona_rid']}")
                    c3.metric("Io (I)", f"{res_nc['io_tot']}/{res_nc['io_rid']}")
                    c4.metric("Destino (D)", f"{res_d['dest_tot']}/{res_d['dest_rid']}")
                    c5.metric("Quintessenza (Q)", f"{q_tot}/{q_rid}")
                    
                    st.write(f"🌟 **Segno Zodiacale**: **{segno_zodiacale}** | **Approdo (Giorno)**: {res_d['approdo']} | **Numero del Sole**: {res_d['sole_tot']}/{res_d['sole_rid']}")
                    st.write(f"📌 **Età dell'Innesto**: **{res_d['eta_innesto']} anni** *(Fase temporale chiave di imprinting e svolta di vita)*")
                    st.write(f"**Origine**: {res_d['origine_rid']} | **Dono**: {res_d['dono_rid']}")
                    st.write(f"**Sfide**: Infanzia {res_d['sfida_infanzia']} | Maturità {res_d['sfida_maturita']} | **Principale**: {res_d['sfida_principale']}")
                    
                    # Dettaglio Piani F.E.M.S.
                    fems = res_nc["fems_count"]
                    tot_lettere = sum(fems.values())
                    st.markdown("#### 🌿 Piani dell'Essere (F.E.M.S.)")
                    pf1, pf2, pf3, pf4 = st.columns(4)
                    pf1.metric("Fisico (4, 8)", f"{fems['Fisico']} lettere")
                    pf2.metric("Emotivo (2,3,6,9)", f"{fems['Emotivo']} lettere")
                    pf3.metric("Mentale (1, 5)", f"{fems['Mentale']} lettere")
                    pf4.metric("Spirituale (7)", f"{fems['Spirituale']} lettere")
                    
                    # Dettaglio Karmico
                    st.markdown("#### ⚖️ Lezioni Karmiche e Archetipi Dominanti")
                    st.write(f"**Lezioni Karmiche (Numeri Mancanti reali)**: {', '.join(map(str, lezioni_karmiche)) if lezioni_karmiche else 'Nessuna'}")
                    st.write(f"**Archetipi Dominanti (Numeri in Eccesso $\ge 3$ volte)**: {', '.join(map(str, archetipi_dominanti)) if archetipi_dominanti else 'Nessuno'}")
                    st.markdown("---")

                    report_matematico = f"""
DATI NUMERICI ESATTI E ANALISI STRUTTURALE:
- Soggetto: {nome_cognome} (Data di nascita: {data_nascita})
- Segno Zodiacale: {segno_zodiacale}
- Anima: {res_nc['anima_tot']}/{res_nc['anima_rid']}
- Persona: {res_nc['persona_tot']}/{res_nc['persona_rid']}
- Io / Espressione: {res_nc['io_tot']}/{res_nc['io_rid']}
- Destino / Cammino: {res_d['dest_tot']}/{res_d['dest_rid']}
- Approdo: {res_d['approdo']}
- Numero del Sole: {res_d['sole_tot']}/{res_d['sole_rid']}
- ETÀ DELL'INNESTO: {res_d['eta_innesto']} anni (Fase temporale di imprinting o evento cruciale, senza interpretazione numerologica dell'età stessa)
- Origine: {res_d['origine_rid']} | Dono: {res_d['dono_rid']}
- Sfide: Infanzia {res_d['sfida_infanzia']}, Maturità {res_d['sfida_maturita']}, Principale {res_d['sfida_principale']}
- Quintessenza: {q_tot}/{q_rid}

CONTEGGIO CORRETTO PIANI DELL'ESSERE (F.E.M.S. via valori pitagorici):
- Piano Fisico (Numeri 4 e 8): {fems['Fisico']} lettere
- Piano Emotivo (Numeri 2, 3, 6, 9): {fems['Emotivo']} lettere
- Piano Mentale (Numeri 1 e 5): {fems['Mentale']} lettere
- Piano Spirituale (Numero 7): {fems['Spirituale']} lettere
(Totale lettere del nome analizzate: {tot_lettere})

LEZIONI KARMICHE E ARCHETIPI DOMINANTI:
- Lezioni Karmiche (Numeri davvero mancanti anche fuori dalla Croce Cardinale): {lezioni_karmiche}
- Archetipi Dominanti (Numeri reiterati 3 o più volte): {archetipi_dominanti}
"""

                    prompt = f"""
Sei un Maestro Numerologo raffinato, colto e saggio della Scuola Numerica.
Scrivi un'analisi narrativa molto ampia, ricca, elegante e priva di qualsiasi errore grammaticale in LINGUA ITALIANA IMPECCABILE.

DATI DI RIFERIMENTO E MATEMATICA:
{report_matematico}

DISPENSA PDF:
{pdf_text}

ISTRUZIONI DI CONTENUTO E REGOLE METODOLOGICHE RIGIDE:
1. CURA FORMALE: Italiano impeccabile. Utilizza esclusivamente i termini corretti ("piano fisico", "piano emotivo", "piano mentale", "piano spirituale"). Non usare mai espressioni scorrette come "pianoforte". Non duplicare parole.
2. SEGNO ZODIACALE: Integra il segno zodiacale ({segno_zodiacale}) all'inizio dell'analisi per arricchire il quadro di partenza dell'anima incarnata.
3. DISTINZIONE SOLE vs INNESTO: 
   - Interpreta il Numero del Sole ({res_d['sole_tot']}/{res_d['sole_rid']}) secondo il suo significato numerologico.
   - Tratta l'Età dell'Innesto ({res_d['eta_innesto']} anni) ESCLUSIVAMENTE come una constatazione temporale di svolta/imprinting vitale, SENZA applicare interpretazioni numerologiche all'età in sé.
4. CAMMINO DEI TALENTI E VOCAZIONE PRATICA: Dedica un'ampia trattazione al Destino ({res_d['dest_tot']}/{res_d['dest_rid']}) e all'Io ({res_nc['io_tot']}/{res_nc['io_rid']}), spiegando in modo dettagliato cosa la persona è portata a fare nella vita pratica e professionale.
5. PIANI DELL'ESSERE (F.E.M.S.): Spiega nel dettaglio la distribuzione dei piani (Fisico={fems['Fisico']}, Emotivo={fems['Emotivo']}, Mentale={fems['Mentale']}, Spirituale={fems['Spirituale']}) in base ai conteggi calcolati.
6. LEZIONI KARMICHE E ARCHETIPI DOMINANTI: Dedica una sezione specifica:
   - Spiega le Lezioni Karmiche ({lezioni_karmiche}) come lezioni o aree di esperienza non ancora integrate.
   - Spiega gli Archetipi Dominanti ({archetipi_dominanti}) come talenti forti o potenziali eccessi energetici.

STRUTTURA DEI CAPITOLI:
- 🌅 **L'Ingresso nel Viaggio, lo Zodiaco e l'Approdo**: Riflessione sull'incarnazione di {nome_cognome} ({data_nascita}), il segno zodiacale ({segno_zodiacale}) e il giorno di Approdo ({res_d['approdo']}).
- 🎶 **La Sinfonia del Nome: Anima, Persona e Io**: Anima ({res_nc['anima_tot']}/{res_nc['anima_rid']}), Persona ({res_nc['persona_tot']}/{res_nc['persona_rid']}) e i mezzi d'azione dell'Io ({res_nc['io_tot']}/{res_nc['io_rid']}).
- 🌱 **Il Cammino dei Talenti e della Professione**: Trattazione ampia delle inclinazioni pratiche del Destino ({res_d['dest_tot']}/{res_d['dest_rid']}), del Numero del Sole ({res_d['sole_tot']}/{res_d['sole_rid']}) e della nota cronologica sull'Età dell'Innesto a {res_d['eta_innesto']} anni.
- 🌿 **I Piani dell'Essere (F.E.M.S.)**: Analisi della struttura energetica tra piano Fisico, Emotivo, Mentale e Spirituale.
- ⚖️ **Lezioni Karmiche e Archetipi Dominanti**: Approfondimento dei numeri mancanti ({lezioni_karmiche}) e dominanti ({archetipi_dominanti}).
- 🌑 **Il Cammino delle Sfide e dell'Ombra**: Le sotto-sfide e la Sfida Principale come semi di superamento.
- 🔮 **La Quintessenza e la Croce Cardinale**: Sintesi d'equilibrio ({q_tot}/{q_rid}) e missione d'Anima.
"""

                    response = client.chat.completions.create(
                        model="gpt-4o-mini",
                        messages=[{"role": "user", "content": prompt}],
                        temperature=0.2
                    )

                    testo_generato = response.choices[0].message.content

                    st.subheader("📖 Profilo Numerologico Narrativo ed Evolutivo")
                    st.markdown(testo_generato)

                    # --- PULSANTE DOWNLOAD WORD ---
                    st.markdown("---")
                    st.subheader("💾 Esporta la Mappa")
                    
                    docx_file = crea_documento_word(nome_cognome, report_matematico, testo_generato)
                    
                    st.download_button(
                        label="📥 Scarica Mappa in formato Word (.docx)",
                        data=docx_file,
                        file_name=f"Mappa_Numerologica_{nome_cognome.replace(' ', '_')}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    )

                except Exception as e:
                    st.error(f"Si è verificato un errore: {e}")
        else:
            st.warning("Compila tutti i campi.")
